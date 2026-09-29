# -*- coding: utf-8 -*-
"""表格解析的行为约束测试

只覆盖「改错了不会报错、只会让表格内容悄悄丢失」的逻辑，不依赖模型与网络。

为什么这几条值得测：
    1. _rows_to_markdown 是 PDF / docx 共用的唯一出口，改错格式两路一起坏。
    2. 不足两行必须返回空串，否则单行表会产出只有表头的残表，污染检索。
    3. docx 建表集成：确认 load_document 真的把表格带出来，而不只是正文。

运行：backend/.venv/Scripts/python -m unittest discover -s backend/tests -v
"""

import sys
import tempfile
import unittest
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from core.loader import _rows_to_markdown, load_document  # noqa: E402


class RowsToMarkdownTest(unittest.TestCase):
    def test_basic_table(self):
        md = _rows_to_markdown([["A", "B"], ["1", "2"]])
        self.assertEqual(md, "| A | B |\n| --- | --- |\n| 1 | 2 |")

    def test_single_row_returns_empty(self):
        self.assertEqual(_rows_to_markdown([["H"]]), "")

    def test_all_empty_rows_returns_empty(self):
        self.assertEqual(_rows_to_markdown([["", ""], ["", ""]]), "")

    def test_ragged_rows_padded(self):
        self.assertIn("| x | y |  |", _rows_to_markdown([["a", "b", "c"], ["x", "y"]]))

    def test_cell_newline_and_pipe_escaped(self):
        md = _rows_to_markdown([["h1", "h2"], ["a\nb", "x|y"]])
        self.assertIn("a b", md)
        self.assertIn("x\\|y", md)


class DocxTableTest(unittest.TestCase):
    def test_docx_table_becomes_markdown(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "s.docx"
            doc = Document()
            doc.add_paragraph("hello-body")
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "PA"
            t.cell(0, 1).text = "PB"
            t.cell(1, 0).text = "v1"
            t.cell(1, 1).text = "v2"
            doc.save(str(p))

            text = load_document(p)
            self.assertIn("hello-body", text)
            self.assertIn("| PA | PB |", text)
            self.assertIn("| v1 | v2 |", text)


if __name__ == "__main__":
    unittest.main()
